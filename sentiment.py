import json
import os
import sys
from datetime import datetime
from pathlib import Path

from azure.ai.textanalytics import TextAnalyticsClient
from azure.core.credentials import AzureKeyCredential
from dotenv import load_dotenv
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

REVIEWS_PATH = Path(__file__).parent / "reviews.json"
OUTPUT_DIR = Path(__file__).parent / "output"

BATCH_SIZE = 10  # Azure sentiment API の1リクエスト最大件数

HEADER_FILL = PatternFill("solid", fgColor="305496")
HEADER_FONT = Font(bold=True, color="FFFFFF")

SENTIMENT_FILLS = {
    "positive": PatternFill("solid", fgColor="C6EFCE"),
    "negative": PatternFill("solid", fgColor="FFC7CE"),
    "mixed":    PatternFill("solid", fgColor="FFEB9C"),
    "neutral":  PatternFill("solid", fgColor="D9D9D9"),
}


def build_client() -> TextAnalyticsClient:
    load_dotenv()
    endpoint = os.environ.get("AZURE_LANGUAGE_ENDPOINT")
    key = os.environ.get("AZURE_LANGUAGE_KEY")
    if not endpoint or not key:
        sys.exit("AZURE_LANGUAGE_ENDPOINT と AZURE_LANGUAGE_KEY を .env に設定してください")
    return TextAnalyticsClient(endpoint=endpoint, credential=AzureKeyCredential(key))


def load_reviews(path: Path) -> tuple[list[str], str]:
    if not path.exists():
        sys.exit(f"レビューファイルが見つかりません: {path}")
    data = json.loads(path.read_text(encoding="utf-8"))
    reviews = data.get("reviews", [])
    if not reviews:
        sys.exit(f"{path} の 'reviews' が空です")
    return reviews, data.get("language", "ja")


def analyze_sentiments(client: TextAnalyticsClient, documents: list[str], language: str) -> list:
    results = []
    for start in range(0, len(documents), BATCH_SIZE):
        chunk = documents[start:start + BATCH_SIZE]
        results.extend(client.analyze_sentiment(
            documents=chunk,
            language=language,
            show_opinion_mining=True,
        ))
    return results


def style_header(sheet) -> None:
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    sheet.freeze_panes = "A2"


def color_sentiment_column(sheet, col_idx: int) -> None:
    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        cell = row[col_idx]
        fill = SENTIMENT_FILLS.get(cell.value)
        if fill:
            cell.fill = fill


def set_widths(sheet, widths: dict[str, int]) -> None:
    for col, width in widths.items():
        sheet.column_dimensions[col].width = width


def write_excel(documents: list[str], results: list, output_path: Path) -> None:
    wb = Workbook()

    # --- summary: 1行 = 1ドキュメント ---
    summary = wb.active
    summary.title = "summary"
    summary.append(["#", "text", "sentiment", "positive", "neutral", "negative", "sentences", "opinions"])
    for i, (text, doc) in enumerate(zip(documents, results), start=1):
        if doc.is_error:
            summary.append([i, text, f"ERROR: {doc.error.code}", None, None, None, None, None])
            continue
        s = doc.confidence_scores
        n_opinions = sum(len(sent.mined_opinions) for sent in doc.sentences)
        summary.append([
            i, text, doc.sentiment,
            round(s.positive, 3), round(s.neutral, 3), round(s.negative, 3),
            len(doc.sentences), n_opinions,
        ])
    style_header(summary)
    color_sentiment_column(summary, col_idx=2)
    set_widths(summary, {"A": 5, "B": 70, "C": 10, "D": 9, "E": 9, "F": 9, "G": 10, "H": 10})
    for row in summary.iter_rows(min_row=2, max_row=summary.max_row):
        row[1].alignment = Alignment(wrap_text=True, vertical="top")

    # --- sentences: 1行 = 1文 ---
    sentences = wb.create_sheet("sentences")
    sentences.append(["doc#", "sent#", "sentiment", "positive", "neutral", "negative", "text"])
    for i, doc in enumerate(results, start=1):
        if doc.is_error:
            continue
        for j, sent in enumerate(doc.sentences, start=1):
            s = sent.confidence_scores
            sentences.append([
                i, j, sent.sentiment,
                round(s.positive, 3), round(s.neutral, 3), round(s.negative, 3),
                sent.text,
            ])
    style_header(sentences)
    color_sentiment_column(sentences, col_idx=2)
    set_widths(sentences, {"A": 6, "B": 6, "C": 10, "D": 9, "E": 9, "F": 9, "G": 70})
    for row in sentences.iter_rows(min_row=2, max_row=sentences.max_row):
        row[6].alignment = Alignment(wrap_text=True, vertical="top")

    # --- opinions: 1行 = (target, assessment) ペア ---
    opinions = wb.create_sheet("opinions")
    opinions.append(["doc#", "sent#", "target", "target_sentiment", "assessment", "assessment_sentiment", "negated"])
    for i, doc in enumerate(results, start=1):
        if doc.is_error:
            continue
        for j, sent in enumerate(doc.sentences, start=1):
            for mined in sent.mined_opinions:
                target = mined.target
                for assessment in mined.assessments:
                    opinions.append([
                        i, j, target.text, target.sentiment,
                        assessment.text, assessment.sentiment, assessment.is_negated,
                    ])
    style_header(opinions)
    color_sentiment_column(opinions, col_idx=3)
    color_sentiment_column(opinions, col_idx=5)
    set_widths(opinions, {"A": 6, "B": 6, "C": 18, "D": 16, "E": 18, "F": 20, "G": 8})

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)


def print_summary(results: list) -> None:
    counts = {"positive": 0, "neutral": 0, "negative": 0, "mixed": 0, "error": 0}
    for doc in results:
        key = "error" if doc.is_error else doc.sentiment
        counts[key] = counts.get(key, 0) + 1
    breakdown = " ".join(f"{k}={v}" for k, v in counts.items() if v > 0)
    print(f"Sentiment breakdown: {breakdown}")


def main() -> None:
    reviews, language = load_reviews(REVIEWS_PATH)
    client = build_client()
    print(f"Analyzing {len(reviews)} reviews (language={language})...")
    results = analyze_sentiments(client, reviews, language)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    output_path = OUTPUT_DIR / f"sentiment_{timestamp}.xlsx"
    write_excel(reviews, results, output_path)

    print_summary(results)
    print(f"Wrote: {output_path}")


if __name__ == "__main__":
    main()
